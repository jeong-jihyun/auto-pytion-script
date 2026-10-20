"""엑셀 요약 보고서 자동 생성 스크립트.

CSV 한 개를 읽어서 엑셀(.xlsx) 보고서를 만듭니다.
- '원본' 시트: CSV 내용 그대로(숫자는 숫자로 저장), 머리글 서식 적용
- '요약' 시트: 그룹별 건수·합계·평균 (엑셀 수식으로 넣어서, 원본을 고치면 같이 바뀝니다)
외부 패키지: openpyxl (requirements.txt 참고)
"""
import argparse
import csv
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def 숫자로(값):
    """문자열이 숫자면 숫자로 바꾸고, 아니면 그대로 돌려줍니다. (1,234 같은 쉼표도 처리)"""
    글자 = str(값).replace(",", "").strip()
    try:
        d = Decimal(글자)
    except InvalidOperation:
        return 값
    return int(d) if d == d.to_integral_value() and "." not in 글자 else float(d)


def 머리글_꾸미기(시트, 열수):
    for c in range(1, 열수 + 1):
        셀 = 시트.cell(row=1, column=c)
        셀.font = Font(bold=True, color="FFFFFF")
        셀.fill = PatternFill("solid", fgColor="305496")
        셀.alignment = Alignment(horizontal="center")


def 열너비_맞추기(시트):
    for 열 in 시트.columns:
        가장긴 = max(len(str(c.value)) if c.value is not None else 0 for c in 열)
        시트.column_dimensions[get_column_letter(열[0].column)].width = min(가장긴 + 4, 40)


def main():
    p = argparse.ArgumentParser(description="CSV를 엑셀 요약 보고서로 만듭니다.")
    p.add_argument("input", help="입력 CSV 파일")
    p.add_argument("-o", "--output", default="output/report.xlsx", help="저장할 xlsx 경로")
    p.add_argument("--group-col", required=True, help="묶을 기준 열 이름 (예: 지점)")
    p.add_argument("--value-col", required=True, help="합계를 낼 숫자 열 이름 (예: 매출)")
    p.add_argument("--encoding", default="utf-8-sig", help="CSV 인코딩 (엑셀에서 저장한 파일이면 cp949)")
    a = p.parse_args()

    with open(a.input, newline="", encoding=a.encoding) as fp:
        읽기 = csv.reader(fp)
        머리글 = next(읽기, None)
        행들 = [r for r in 읽기 if any(x.strip() for x in r)]
    if not 머리글:
        print("CSV가 비어 있습니다.")
        return 1
    for 이름 in (a.group_col, a.value_col):
        if 이름 not in 머리글:
            print(f"열 '{이름}' 이(가) 없습니다. 있는 열: {머리글}")
            return 1
    g = 머리글.index(a.group_col) + 1  # 엑셀 열 번호는 1부터
    v = 머리글.index(a.value_col) + 1
    gL, vL = get_column_letter(g), get_column_letter(v)

    wb = Workbook()
    원본 = wb.active
    원본.title = "원본"
    원본.append(머리글)
    for r in 행들:
        원본.append([숫자로(x) for x in r])
    머리글_꾸미기(원본, len(머리글))
    원본.freeze_panes = "A2"
    열너비_맞추기(원본)

    끝 = len(행들) + 1
    그룹들 = []
    for r in 행들:
        if r[g - 1] not in 그룹들:
            그룹들.append(r[g - 1])
    요약 = wb.create_sheet("요약")
    요약.append([a.group_col, "건수", f"{a.value_col} 합계", f"{a.value_col} 평균"])
    for i, 이름 in enumerate(그룹들, start=2):
        범위g = f"원본!${gL}$2:${gL}${끝}"
        범위v = f"원본!${vL}$2:${vL}${끝}"
        요약.append([이름, f"=COUNTIF({범위g},A{i})", f"=SUMIF({범위g},A{i},{범위v})",
                   f"=IF(B{i}=0,0,C{i}/B{i})"])
    합계행 = len(그룹들) + 2
    요약.append(["전체", f"=SUM(B2:B{합계행 - 1})", f"=SUM(C2:C{합계행 - 1})",
               f"=IF(B{합계행}=0,0,C{합계행}/B{합계행})"])
    요약.cell(row=합계행, column=1).font = Font(bold=True)
    for r in range(2, 합계행 + 1):
        요약.cell(row=r, column=3).number_format = "#,##0"
        요약.cell(row=r, column=4).number_format = "#,##0.0"
    머리글_꾸미기(요약, 4)
    열너비_맞추기(요약)

    out = Path(a.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    # 엑셀을 열지 않아도 결과를 확인할 수 있게, 파이썬으로 계산한 합계도 화면에 보여 줍니다.
    합 = {}
    for r in 행들:
        x = 숫자로(r[v - 1])
        합[r[g - 1]] = 합.get(r[g - 1], 0) + (x if isinstance(x, (int, float)) else 0)
    print(f"완료: 원본 {len(행들)}행, 그룹 {len(그룹들)}개 -> {out}")
    for k, s in 합.items():
        print(f"  {k}: 합계 {s:,}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
