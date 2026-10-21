"""엑셀(xlsx)을 CSV로 바꾸는 스크립트.

엑셀 파일의 시트마다 CSV 파일을 하나씩 만듭니다.
수식 칸은 계산된 값(마지막으로 저장된 값)을 씁니다.
필요한 패키지: openpyxl (pip install openpyxl)
"""
import argparse
import csv
import re
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    raise SystemExit("openpyxl 이 필요합니다. 'pip install openpyxl' 로 설치하세요.")


def cell_text(v):
    """칸 값을 글자로 바꿉니다. 비어 있으면 빈 글자."""
    if v is None:
        return ""
    if hasattr(v, "isoformat"):  # 날짜·시간
        if hasattr(v, "hour"):  # 시각이 00:00:00 이면 날짜만 씀
            return v.date().isoformat() if v.time().isoformat() == "00:00:00" else v.isoformat(sep=" ")
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def main():
    p = argparse.ArgumentParser(description="xlsx -> CSV (시트별)")
    p.add_argument("input", type=Path)
    p.add_argument("-o", "--output", type=Path, default=Path("output"))
    p.add_argument("--sheet", help="이 이름의 시트만 변환 (생략하면 모두)")
    a = p.parse_args()

    wb = load_workbook(a.input, data_only=True, read_only=True)
    names = [a.sheet] if a.sheet else wb.sheetnames
    for n in names:
        if n not in wb.sheetnames:
            raise SystemExit(f"오류: '{n}' 시트가 없습니다. 시트 목록: {wb.sheetnames}")
    a.output.mkdir(parents=True, exist_ok=True)
    for n in names:
        rows = [[cell_text(c) for c in row] for row in wb[n].iter_rows(values_only=True)]
        # 맨 아래의 완전히 빈 줄은 버립니다
        while rows and not any(rows[-1]):
            rows.pop()
        out = a.output / (re.sub(r'[\\/:*?"<>|]', "_", n) + ".csv")
        with out.open("w", newline="", encoding="utf-8-sig") as fp:
            csv.writer(fp).writerows(rows)
        print(f"시트 '{n}': {len(rows)}줄 -> {out}")


if __name__ == "__main__":
    main()
