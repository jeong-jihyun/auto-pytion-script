"""거래처별 청구서(docx) 자동 생성 스크립트.

품목 CSV(거래처, 품목, 수량, 단가)를 읽어 거래처마다 워드(docx) 청구서를 한 장씩 만듭니다.
공급가액, 부가세(10%), 합계를 계산해 표 아래에 적습니다.
외부 패키지 python-docx 가 필요합니다. (pip install -r requirements.txt)
"""
import argparse
import csv
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def safe_name(text):
    """파일 이름에 쓸 수 없는 글자를 _ 로 바꿉니다."""
    return re.sub(r'[\\/:*?"<>|\s]+', "_", text)


def make_invoice(client, items, issue_date, path):
    doc = Document()
    title = doc.add_heading("청 구 서", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"수신: {client} 귀하")
    doc.add_paragraph(f"발행일: {issue_date}")

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, head in zip(table.rows[0].cells, ["품목", "수량", "단가", "금액", "비고"]):
        cell.text = head
    supply = 0
    for it in items:
        amount = it["수량"] * it["단가"]
        supply += amount
        cells = table.add_row().cells
        cells[0].text = it["품목"]
        cells[1].text = f"{it['수량']:,}"
        cells[2].text = f"{it['단가']:,}"
        cells[3].text = f"{amount:,}"
        cells[4].text = ""
    vat = supply // 10  # 부가세 10%, 원 미만 버림
    doc.add_paragraph()
    doc.add_paragraph(f"공급가액: {supply:,}원")
    doc.add_paragraph(f"부가세(10%): {vat:,}원")
    doc.add_paragraph(f"합계: {supply + vat:,}원")
    doc.save(path)
    return supply, vat


def main():
    p = argparse.ArgumentParser(description="거래처별 청구서 docx 만들기")
    p.add_argument("input", type=Path, help="품목 CSV (열: 거래처,품목,수량,단가)")
    p.add_argument("-o", "--outdir", type=Path, default=Path("output"))
    p.add_argument("--date", default=date.today().isoformat(), help="발행일(기본: 오늘)")
    a = p.parse_args()

    groups = defaultdict(list)
    with a.input.open(newline="", encoding="utf-8-sig") as fp:
        for n, row in enumerate(csv.DictReader(fp), start=2):
            try:
                groups[row["거래처"].strip()].append({
                    "품목": row["품목"].strip(),
                    "수량": int(row["수량"].replace(",", "")),
                    "단가": int(row["단가"].replace(",", "")),
                })
            except (ValueError, KeyError, AttributeError):
                print(f"[건너뜀] {n}번째 줄을 읽을 수 없습니다: {row}")
    a.outdir.mkdir(parents=True, exist_ok=True)
    for client, items in groups.items():
        path = a.outdir / f"청구서_{safe_name(client)}.docx"
        supply, vat = make_invoice(client, items, a.date, path)
        print(f"{client}: 품목 {len(items)}개, 합계 {supply + vat:,}원 -> {path}")
    print(f"청구서 {len(groups)}장을 만들었습니다.")


if __name__ == "__main__":
    main()
