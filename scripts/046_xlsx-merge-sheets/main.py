"""여러 엑셀(xlsx) 파일을 한 파일의 '시트'로 모아 주는 스크립트.

폴더의 xlsx 파일마다 (파일 안의 첫 번째 시트를) 새 파일의 시트 하나로 복사합니다.
시트 이름은 원래 파일 이름이 됩니다(엑셀 규칙상 31자 이내, 특수문자 제외).
또 맨 앞에 '목차' 시트를 만들어 어떤 파일이 몇 행인지 적어 둡니다.
값만 복사합니다(수식은 엑셀이 마지막에 저장해 둔 계산 결과 값으로 들어갑니다).
"""
import argparse
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("openpyxl 이 필요합니다:  pip install openpyxl")
    sys.exit(1)


def sheet_name(stem, used):
    """엑셀 시트 이름 규칙(31자, []:*?/\\ 금지)에 맞게 바꾸고 중복을 피합니다."""
    name = re.sub(r"[\[\]:*?/\\]", "_", stem)[:31] or "sheet"
    base, i = name, 2
    while name in used:
        suffix = f"_{i}"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name)
    return name


def main():
    p = argparse.ArgumentParser(description="여러 xlsx 를 한 파일의 시트로 모으기")
    p.add_argument("folder", type=Path, help="xlsx 파일이 있는 폴더")
    p.add_argument("-o", "--output", type=Path, default=Path("output/merged.xlsx"))
    a = p.parse_args()

    files = sorted(f for f in a.folder.glob("*.xlsx") if not f.name.startswith("~$"))
    if not files:
        print("xlsx 파일이 없습니다.")
        sys.exit(1)
    out = Workbook()
    toc = out.active
    toc.title = "목차"
    toc.append(["원본 파일", "시트 이름", "행 수", "열 수"])
    used = {"목차"}
    for f in files:
        wb = load_workbook(f, data_only=True, read_only=True)
        src = wb.worksheets[0]
        name = sheet_name(f.stem, used)
        ws = out.create_sheet(name)
        rows = 0
        for row in src.iter_rows(values_only=True):
            ws.append(list(row))
            rows += 1
        toc.append([f.name, name, rows, src.max_column])
        print(f"{f.name}: {rows}행 -> 시트 '{name}'")
        wb.close()
    a.output.parent.mkdir(parents=True, exist_ok=True)
    out.save(a.output)
    print(f"저장: {a.output} (시트 {len(files)}개 + 목차)")


if __name__ == "__main__":
    main()
