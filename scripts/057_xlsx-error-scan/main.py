"""엑셀(xlsx) 파일 속 '문제 칸'을 찾아 알려 주는 스크립트. 파일은 읽기만 합니다.

찾는 것
1) 오류 값 칸: #REF!, #DIV/0!, #N/A, #VALUE!, #NAME?, #NUM!, #NULL!
2) 수식 안에 #REF! 가 들어 있는 칸 (지워진 셀을 가리키는 깨진 수식)
3) 숫자처럼 보이지만 글자로 저장된 칸 (예: '1,234' 문자열) -> 합계에서 빠질 수 있음
4) 다른 파일을 가리키는 수식(외부 링크)
계산 결과(캐시 값)는 엑셀에서 마지막으로 저장될 때의 값입니다.
"""
import argparse
import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ERRORS = {"#REF!", "#DIV/0!", "#N/A", "#VALUE!", "#NAME?", "#NUM!", "#NULL!"}
NUM_TEXT = re.compile(r"^\s*[-+]?\d{1,3}(,\d{3})+(\.\d+)?\s*$|^\s*[-+]?\d+(\.\d+)?\s*$")


def main():
    p = argparse.ArgumentParser(description="엑셀 문제 칸 찾기")
    p.add_argument("input", type=Path, help="xlsx 파일")
    p.add_argument("-o", "--output", type=Path, default=Path("output/xlsx_issues.csv"))
    a = p.parse_args()

    try:
        wb_f = load_workbook(a.input, data_only=False)    # 수식 그대로
        wb_v = load_workbook(a.input, data_only=True)     # 저장된 계산 결과
    except Exception as exc:
        print(f"[오류] 파일을 열 수 없습니다: {exc}")
        sys.exit(1)

    issues = []
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                cached = wv[c.coordinate].value
                if isinstance(v, str) and v.startswith("="):
                    if "#REF!" in v:
                        issues.append((ws.title, c.coordinate, "깨진 수식(#REF!)", v))
                    if re.search(r"\[\d+\]|'?\w:\\|https?://", v):
                        issues.append((ws.title, c.coordinate, "외부 파일 링크 수식", v))
                    if isinstance(cached, str) and cached in ERRORS:
                        issues.append((ws.title, c.coordinate, f"오류 값 {cached}", v))
                elif isinstance(v, str) and v.strip() in ERRORS:
                    issues.append((ws.title, c.coordinate, f"오류 값 {v.strip()}", v))
                elif isinstance(v, str) and NUM_TEXT.match(v) and c.data_type == "s":
                    issues.append((ws.title, c.coordinate, "숫자인데 글자로 저장됨", v))

    a.output.parent.mkdir(parents=True, exist_ok=True)
    with a.output.open("w", newline="", encoding="utf-8-sig") as fp:
        w = csv.writer(fp)
        w.writerow(["시트", "칸", "문제", "내용"])
        w.writerows(issues)
    print(f"[완료] 시트 {len(wb_f.worksheets)}개 검사, 문제 칸 {len(issues)}개 -> {a.output}")
    for sh, cell, why, val in issues:
        print(f"  - {sh}!{cell}: {why} ({val})")


if __name__ == "__main__":
    main()
