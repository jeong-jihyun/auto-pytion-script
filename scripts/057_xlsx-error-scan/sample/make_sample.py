"""샘플 xlsx(book.xlsx)를 만드는 보조 스크립트. 한 번만 실행하면 됩니다."""
from pathlib import Path

from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "매출"
ws.append(["항목", "금액"])
ws.append(["A", 1000])
ws.append(["B", "2,500"])        # 글자로 저장된 숫자
ws.append(["C", "#DIV/0!"])      # 오류 값
ws.append(["합계", "=SUM(B2:B4)"])
ws["D2"] = "=SUM(#REF!)"         # 깨진 수식
ws["D3"] = "='[2]Sheet1'!A1"     # 외부 링크 수식
wb.save(Path(__file__).with_name("book.xlsx"))
print("book.xlsx 생성")
