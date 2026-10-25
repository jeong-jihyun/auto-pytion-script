"""간단한 마크다운(.md) 문서를 워드(.docx)로 바꾸는 스크립트. 주간보고·회의록을 워드로 낼 때 씁니다.

지원: # ~ ### 제목, - 또는 * 글머리표, 1. 번호 목록, **굵게**, 마크다운 표(| 로 구분), 일반 문단.
외부 패키지 python-docx가 필요합니다. (pip install -r requirements.txt)
"""
import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("[오류] python-docx가 필요합니다: pip install python-docx")
    sys.exit(1)


def add_runs(par, text):
    """**굵게** 표시를 워드의 굵은 글씨로 바꿔 문단에 넣습니다."""
    for i, piece in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if piece:
            par.add_run(piece).bold = (i % 2 == 1)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main():
    p = argparse.ArgumentParser(description="마크다운 → 워드(docx)")
    p.add_argument("input", type=Path)
    p.add_argument("-o", "--output", type=Path, default=Path("output/report.docx"))
    a = p.parse_args()
    if not a.input.is_file():
        print(f"[오류] 파일이 없습니다: {a.input}")
        sys.exit(1)

    lines = a.input.read_text(encoding="utf-8-sig").splitlines()
    doc = Document()
    i, counts = 0, {"제목": 0, "문단": 0, "목록": 0, "표": 0}
    while i < len(lines):
        s = lines[i].rstrip()
        if not s.strip():
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", s)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
            counts["제목"] += 1
        elif s.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [split_row(x) for x in block if not re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", x)]
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = "Table Grid"
            for ri, r in enumerate(rows):
                for ci in range(ncol):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], r[ci] if ci < len(r) else "")
                    if ri == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            counts["표"] += 1
            continue
        elif re.match(r"^\s*[-*]\s+", s):
            add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", s))
            counts["목록"] += 1
        elif re.match(r"^\s*\d+\.\s+", s):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s+", "", s))
            counts["목록"] += 1
        else:
            add_runs(doc.add_paragraph(), s)
            counts["문단"] += 1
        i += 1

    a.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(a.output)
    print("변환 완료:", ", ".join(f"{k} {v}" for k, v in counts.items()), f"→ {a.output}")


if __name__ == "__main__":
    main()
