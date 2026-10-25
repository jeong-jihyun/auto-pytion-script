"""두 엑셀 파일(이전본·수정본)을 셀 하나하나 비교해 바뀐 곳을 찾는 스크립트.

생각의 흐름:
  1) 두 파일에서 같은 이름의 시트끼리 짝을 짓습니다. 한쪽에만 있는 시트도 알려 줍니다.
  2) 두 시트의 사용 범위 중 큰 쪽까지 모든 셀을 돌며 값을 비교합니다.
     - 숫자는 아주 작은 차이(부동소수 오차)는 같다고 봅니다(--tol).
     - 수식은 기본으로 '수식 글자'를 비교합니다(계산 결과는 엑셀이 저장한 값이 없을 수 있음).
  3) 바뀐 셀 목록 CSV + 수정본 사본에 바뀐 셀을 노란색으로 칠한 xlsx를 만듭니다.
     원본 파일은 건드리지 않습니다.
"""
import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

YELLOW = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
GREEN = PatternFill("solid", start_color="D9EAD3", end_color="D9EAD3")


def same(a, b, tol):
    if a in (None, "") and b in (None, ""):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= tol
    return a == b


def main():
    p = argparse.ArgumentParser(description="엑셀 두 파일 셀 단위 비교")
    p.add_argument("old", type=Path, help="이전 xlsx")
    p.add_argument("new", type=Path, help="수정한 xlsx")
    p.add_argument("--tol", type=float, default=1e-9, help="숫자를 같다고 볼 오차")
    p.add_argument("--values", action="store_true", help="수식 대신 저장된 계산값으로 비교")
    p.add_argument("-o", "--outdir", type=Path, default=Path("output"))
    a = p.parse_args()

    wb_old = load_workbook(a.old, data_only=a.values)
    wb_new = load_workbook(a.new, data_only=a.values)
    mark = load_workbook(a.new)  # 색칠용 사본(수식 유지)
    changes = []
    only_old = [s for s in wb_old.sheetnames if s not in wb_new.sheetnames]
    only_new = [s for s in wb_new.sheetnames if s not in wb_old.sheetnames]
    for name in wb_new.sheetnames:
        if name not in wb_old.sheetnames:
            continue
        so, sn, sm = wb_old[name], wb_new[name], mark[name]
        rows = max(so.max_row, sn.max_row)
        cols = max(so.max_column, sn.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                vo, vn = so.cell(r, c).value, sn.cell(r, c).value
                if same(vo, vn, a.tol):
                    continue
                kind = "추가" if vo in (None, "") else ("삭제" if vn in (None, "") else "변경")
                addr = f"{get_column_letter(c)}{r}"
                changes.append([name, addr, kind, "" if vo is None else vo, "" if vn is None else vn])
                cell = sm.cell(r, c)
                cell.fill = GREEN if kind == "추가" else YELLOW
                cell.comment = Comment(f"이전 값: {vo}", "비교도구")

    a.outdir.mkdir(parents=True, exist_ok=True)
    out_csv = a.outdir / "changes.csv"
    with out_csv.open("w", newline="", encoding="utf-8-sig") as fp:
        w = csv.writer(fp)
        w.writerow(["시트", "셀", "구분", "이전값", "새값"])
        w.writerows(changes)
    out_xlsx = a.outdir / f"{a.new.stem}_비교표시.xlsx"
    mark.save(out_xlsx)

    print(f"바뀐 셀 {len(changes)}개 → {out_csv}, 색칠본 {out_xlsx}")
    if only_old:
        print(f"  ! 이전 파일에만 있는 시트: {', '.join(only_old)}")
    if only_new:
        print(f"  ! 수정 파일에만 있는 시트: {', '.join(only_new)}")
    for s, addr, k, vo, vn in changes[:15]:
        print(f"  - [{s}] {addr} {k}: {vo if vo != '' else '(빈칸)'} → {vn if vn != '' else '(빈칸)'}")


if __name__ == "__main__":
    main()
