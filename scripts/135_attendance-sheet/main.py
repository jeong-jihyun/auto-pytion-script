"""참석자 명단 CSV로 '교육·회의 출석부(서명부)' 엑셀 파일을 만드는 스크립트.

생각의 흐름:
  1) 교육·행사 때마다 엑셀에서 명단을 붙여 넣고 서명란·테두리·인쇄 설정을 손으로 맞추느라 시간이 듭니다.
  2) 명단 CSV(이름, 부서 등)만 주면 다음을 자동으로 만듭니다.
       - 위쪽: 행사명, 일시, 장소
       - 표: 번호 + 명단 열 + 서명 칸(--sign-cols 개수만큼, 예: 오전/오후)
       - 아래: 총원·참석·불참 칸(참석 수는 서명란이 빈칸 아닌 사람 수를 세는 COUNTA 수식)
  3) --group 열(예: 부서)을 주면 그 값별로 시트를 나눠, 부서별로 돌릴 수 있게 합니다.
  4) 인쇄를 고려해 A4 세로, 가로 1쪽 맞춤, 머리글 줄 반복 인쇄를 설정합니다.
  5) 개인정보를 줄이기 위해 --cols 로 필요한 열만 고를 수 있게 합니다(전화번호 등 제외).
"""
import argparse
import csv
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def safe_title(s):
    for ch in '[]:*?/\\':
        s = s.replace(ch, "_")
    return (s or "명단")[:31]


def build_sheet(ws, people, cols, sign_cols, title, when, where):
    ncol = 1 + len(cols) + len(sign_cols)
    last = get_column_letter(ncol)
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = f"일시: {when}    장소: {where}"
    ws["A2"].alignment = Alignment(horizontal="center")

    head = ["번호"] + cols + sign_cols
    hr = 4
    for i, h in enumerate(head, start=1):
        c = ws.cell(row=hr, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    for n, person in enumerate(people, start=1):
        r = hr + n
        ws.cell(row=r, column=1, value=n)
        for j, col in enumerate(cols, start=2):
            ws.cell(row=r, column=j, value=person.get(col, ""))
        for j in range(1, ncol + 1):
            cell = ws.cell(row=r, column=j)
            cell.border = BOX
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 28   # 서명하기 좋게 줄 높이를 넉넉히

    end = hr + len(people)
    s = end + 2
    ws.cell(row=s, column=1, value="총원")
    ws.cell(row=s, column=2, value=len(people))
    first_sign = get_column_letter(2 + len(cols))
    ws.cell(row=s + 1, column=1, value="참석")
    # 첫 번째 서명란에 무언가 적힌 사람 수를 셉니다(서명 후 엑셀에 O 표시를 하면 자동 계산)
    ws.cell(row=s + 1, column=2, value=f"=COUNTA({first_sign}{hr + 1}:{first_sign}{max(end, hr + 1)})")
    ws.cell(row=s + 2, column=1, value="불참")
    ws.cell(row=s + 2, column=2, value=f"=B{s}-B{s + 1}")

    ws.column_dimensions["A"].width = 6
    for j in range(2, ncol + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16 if j <= 1 + len(cols) else 20
    ws.print_title_rows = f"{hr}:{hr}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    p = argparse.ArgumentParser(description="출석부(서명부) 엑셀 만들기")
    p.add_argument("roster", type=Path, help="명단 CSV")
    p.add_argument("--title", default="교육 출석부")
    p.add_argument("--when", default="", help="일시 문구")
    p.add_argument("--where", default="", help="장소 문구")
    p.add_argument("--cols", default="", help="표에 넣을 열, 쉼표로(기본: 모든 열)")
    p.add_argument("--sign-cols", default="서명", help="서명 칸 이름들, 쉼표로. 예: 오전서명,오후서명")
    p.add_argument("--group", default="", help="이 열 값별로 시트를 나눔(예: 부서)")
    p.add_argument("--sort", default="", help="정렬 기준 열(예: 이름)")
    p.add_argument("-o", "--output", type=Path, default=Path("output/attendance.xlsx"))
    a = p.parse_args()

    with a.roster.open(newline="", encoding="utf-8-sig") as fp:
        reader = csv.DictReader(fp)
        header = reader.fieldnames or []
        people = [{k: (v or "").strip() for k, v in r.items()} for r in reader]
    cols = [c.strip() for c in a.cols.split(",") if c.strip()] or header
    for c in cols + ([a.group] if a.group else []) + ([a.sort] if a.sort else []):
        if c not in header:
            raise SystemExit(f"[오류] 명단에 없는 열: {c} (있는 열: {', '.join(header)})")
    sign_cols = [c.strip() for c in a.sign_cols.split(",") if c.strip()]
    if a.sort:
        people.sort(key=lambda r: r[a.sort])

    groups = OrderedDict()
    if a.group:
        for r in people:
            groups.setdefault(r[a.group] or "미지정", []).append(r)
    else:
        groups["출석부"] = people

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in groups.items():
        ws = wb.create_sheet(safe_title(name))
        t = f"{a.title} ({name})" if a.group else a.title
        build_sheet(ws, rows, cols, sign_cols, t, a.when, a.where)
        print(f"  시트 '{ws.title}': {len(rows)}명")
    a.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.output)
    print(f"총 {len(people)}명, 시트 {len(groups)}개 → 저장: {a.output}")


if __name__ == "__main__":
    main()
